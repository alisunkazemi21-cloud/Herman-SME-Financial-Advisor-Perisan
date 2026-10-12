"""Temporary Quick intake: real local parsers, synthetic documents, no OCR quality claim."""

# ruff: noqa: E402
import asyncio
import hashlib
import io
import json
import os
import subprocess
import sys
import threading
import time
import zipfile

import openpyxl
import pytest
from PIL import Image

pytest.importorskip("psycopg")
pytest.importorskip("fastapi")
from fastapi.testclient import TestClient
from starlette.requests import Request
from test_backend import backend as backend_fixture
from test_backend import database as database_fixture

from src.backend import process_guard, quick_document_api, quick_documents, quick_extract
from src.backend.advisor import QuickContextInput, quick_context
from src.backend.api import create_app
from src.backend.database import AccessDenied
from src.backend.process_guard import run_contained
from src.backend.quick_document_api import receive_and_extract
from src.backend.quick_document_models import MAX_UPLOAD, DocumentError, ExtractedUnit
from src.backend.quick_documents import extract_temporary, pack_sources
from src.backend.quick_financial import QuickSource

backend, database = backend_fixture, database_fixture
CSV = "شرح,مبلغ\nخرید گوشت,۱۲۳۴\n".encode("utf-8")


def workbook_bytes(formula=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "فروش"
    sheet.append(["شرح", "مبلغ"])
    sheet.append(["قهوه", "=1+2" if formula else "۱۲۳"])
    book.create_sheet("خرید").append(["گوشت", 456])
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def pdf_bytes(pages=1, text="Invoice 123", width=300, height=200):
    """Minimal independent PDF fixture; no extra document-generation dependency."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    kids = []
    for _ in range(pages):
        page_id = len(objects) + 1
        kids.append(f"{page_id} 0 R")
        stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode("ascii")
        objects += [
            (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] "
             f"/Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> "
             f"/Contents {page_id + 1} 0 R >>").encode(),
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream",
        ]
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {pages} >>".encode()
    output, offsets = bytearray(b"%PDF-1.4\n"), [0]
    for i, value in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{i} 0 obj\n".encode() + value + b"\nendobj\n")
    start = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode())
    return bytes(output)


def pieces(result):
    found = {}
    for source in result["sources"]:
        for location in source["origin"]["locations"]:
            found.setdefault(location["locator"], []).append(
                (location["unit_offset"], source["text_fa"][location["start"]:location["end"]])
            )
    return {locator: "".join(text for _, text in sorted(parts)) for locator, parts in found.items()}


@pytest.mark.parametrize("format,content,expected", [
    ("csv", CSV, {"CSV:row:1": '["شرح", "مبلغ"]', "CSV:row:2": '["خرید گوشت", "۱۲۳۴"]'}),
    ("xlsx", workbook_bytes(), {"فروش:row:1": '["شرح", "مبلغ"]', "فروش:row:2": '["قهوه", "۱۲۳"]',
                               "خرید:row:1": '["گوشت", "456"]'}),
    ("pdf", pdf_bytes(2), {"page:1": "Invoice 123", "page:2": "Invoice 123"}),
], ids=["csv", "xlsx", "pdf"])
def test_real_parser_subprocess_preserves_all_locations_and_cleans_files(tmp_path, format, content, expected):
    result = extract_temporary(content, format, temporary_parent=tmp_path)
    assert pieces(result) == expected
    assert result["document"]["sha256"] == hashlib.sha256(content).hexdigest()
    assert not result["persisted"] and not result["history_used"]
    assert result["verification"] == "locally_extracted_unverified"
    assert result["dependencies"]["pydantic"]
    assert list(tmp_path.iterdir()) == []
    assert "source_file" not in json.dumps(result) and str(tmp_path) not in json.dumps(result)
    for source in result["sources"]:
        QuickSource.model_validate(source)
        assert source["origin"]["file_sha256"] == result["document"]["sha256"]


def test_long_units_round_trip_across_chunks_with_exact_offsets():
    units = [ExtractedUnit(locator="page:1", text="الف " * 3600, method="pdf_text", confidence=1),
             ExtractedUnit(locator="page:2", text="پایان", method="pdf_text", confidence=1)]
    sources = pack_sources(units, "a" * 64, "pdf")
    assert len(sources) == 3 and all(len(source["text_fa"]) <= 6000 for source in sources)
    assert pieces({"sources": sources}) == {unit.locator: unit.text for unit in units}
    assert sources[1]["origin"]["locations"][0]["unit_offset"] == 6000
    context = quick_context(QuickContextInput(question_fa="این سند را توضیح بده", sources=sources))
    assert context["financial"] is None and context["inventory"] is None
    assert all(source["verification"] == "user_supplied_unverified" for source in context["supplied_sources"])


@pytest.mark.parametrize("change", ["hash", "end", "overlap", "path", "format"])
def test_copied_origin_is_bounded_and_text_consistent(change):
    unit = ExtractedUnit(locator="page:1", text="مبلغ ۱۲۳", method="pdf_text", confidence=1)
    source = pack_sources([unit], "a" * 64, "pdf")[0]
    if change == "hash":
        source["text_fa"] = "متن عوض شده"
    elif change == "end":
        source["origin"]["locations"][0]["end"] = 5999
    elif change == "overlap":
        source["origin"]["locations"] *= 2
    elif change == "path":
        source["origin"]["source_file"] = "private/file"
    else:
        source["origin"]["format"] = "../csv"
    with pytest.raises(ValueError):
        QuickSource.model_validate(source)


@pytest.mark.parametrize("format,content,code", [
    ("csv", b"\xff", "invalid_document"), ("xlsx", b"not a zip", "invalid_document"),
    ("csv", b"a\n=1+2\n", "unsupported_document"),
    ("xlsx", workbook_bytes(True), "unsupported_document"),
    ("csv", b"\n\n", "no_text"), ("pdf", pdf_bytes(9), "resource_limit"),
    ("pdf", pdf_bytes(width=100000, height=100000), "resource_limit"),
], ids=["bad-utf8", "bad-zip", "csv-formula", "xlsx-formula", "blank", "page-limit", "raster-limit"])
def test_real_parser_failures_are_generic_and_remove_temporary_files(tmp_path, format, content, code):
    with pytest.raises(DocumentError) as caught:
        extract_temporary(content, format, temporary_parent=tmp_path)
    assert caught.value.code == code
    assert list(tmp_path.iterdir()) == []
    assert str(tmp_path) not in str(caught.value)


@pytest.mark.parametrize("suffix,image_format", [("png", "PNG"), ("jpg", "JPEG"), ("jpeg", "JPEG"),
                                                   ("tiff", "TIFF"), ("bmp", "BMP")])
def test_image_adapters_preserve_unverified_ocr_units(tmp_path, monkeypatch, suffix, image_format):
    source = tmp_path / ("source." + suffix)
    Image.new("RGB", (24, 24)).save(source, format=image_format)
    monkeypatch.setattr(quick_extract.shutil, "which", lambda _: "synthetic-engine")
    monkeypatch.setattr("pytesseract.get_languages", lambda: ["fas", "eng"])
    monkeypatch.setattr(quick_extract.PersianFinancialExtractor, "recognize", lambda *_: ("مبلغ ۱۲۳", 0.75))
    units = quick_extract.extract_units(source, "tesseract")
    assert len(units) == 1 and units[0].locator == "image:1"
    assert units[0].text == "مبلغ ۱۲۳" and units[0].confidence == 0.75


def test_blank_pdf_uses_local_ocr_with_page_reference(tmp_path, monkeypatch):
    path = tmp_path / "source.pdf"
    path.write_bytes(pdf_bytes(text=""))
    monkeypatch.setattr(quick_extract.shutil, "which", lambda _: "synthetic-engine")
    monkeypatch.setattr("pytesseract.get_languages", lambda: ["fas", "eng"])
    monkeypatch.setattr(quick_extract.PersianFinancialExtractor, "recognize", lambda *_: ("رسید فارسی", 0.6))
    units = quick_extract.extract_units(path, "tesseract")
    assert units[0].locator == "page:1" and units[0].method == "tesseract"
    assert units[0].text == "رسید فارسی" and units[0].confidence == 0.6


def test_multiframe_image_rejected_before_ocr(tmp_path, monkeypatch):
    path = tmp_path / "source.tiff"
    image = Image.new("RGB", (5, 5))
    image.save(path, save_all=True, append_images=[image])
    monkeypatch.setattr(quick_extract.PersianFinancialExtractor, "recognize", lambda *_: pytest.fail("no OCR"))
    with pytest.raises(DocumentError, match="unsupported_document"):
        quick_extract.extract_units(path, "tesseract")


def test_missing_local_engine_is_explicit(tmp_path, monkeypatch):
    path = tmp_path / "source.png"
    Image.new("RGB", (5, 5)).save(path)
    monkeypatch.setattr(quick_extract.shutil, "which", lambda _: None)
    with pytest.raises(DocumentError, match="engine_unavailable"):
        quick_extract.extract_units(path, "tesseract")


def test_missing_persian_language_pack_is_explicit(tmp_path, monkeypatch):
    path = tmp_path / "source.png"
    Image.new("RGB", (5, 5)).save(path)
    monkeypatch.setattr(quick_extract.shutil, "which", lambda _: "synthetic-engine")
    monkeypatch.setattr("pytesseract.get_languages", lambda: ["eng"])
    monkeypatch.setattr(quick_extract.PersianFinancialExtractor, "recognize", lambda *_: pytest.fail("no OCR"))
    with pytest.raises(DocumentError, match="engine_unavailable"):
        quick_extract.extract_units(path, "tesseract")


def test_missing_easyocr_models_do_not_create_external_directories(tmp_path, monkeypatch):
    path = tmp_path / "source.png"
    Image.new("RGB", (5, 5)).save(path)
    models = tmp_path / "not-installed"
    monkeypatch.setattr(quick_extract.PersianFinancialExtractor, "recognize", lambda *_: pytest.fail("no OCR"))
    with pytest.raises(DocumentError, match="engine_unavailable"):
        quick_extract.extract_units(path, "easyocr", str(models))
    assert not models.exists()


@pytest.mark.parametrize("content", [b"a\n" * 2001, ("," * 50).encode(), b"a" * 6001])
def test_table_resource_bounds(tmp_path, content):
    path = tmp_path / "source.csv"
    path.write_bytes(content)
    with pytest.raises(DocumentError, match="resource_limit"):
        quick_extract.extract_units(path, "tesseract")


def test_undersized_xlsx_dimensions_do_not_omit_rows(tmp_path):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(workbook_bytes())) as source, zipfile.ZipFile(output, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<dimension ref="A1:B2"', b'<dimension ref="A1:A1"')
            target.writestr(name, data)
    path = tmp_path / "source.xlsx"
    path.write_bytes(output.getvalue())
    units = quick_extract.extract_units(path, "tesseract")
    assert units[1].locator == "فروش:row:2" and "۱۲۳" in units[1].text


def test_sheet_name_whitespace_remains_distinct_in_source_references(tmp_path):
    book = openpyxl.Workbook()
    book.active.title = " فروش"
    book.active.append(["۱۲۳"])
    book.create_sheet("فروش").append(["۴۵۶"])
    path = tmp_path / "source.xlsx"
    book.save(path)
    book.close()
    units = quick_extract.extract_units(path, "tesseract")
    sources = pack_sources(units, "c" * 64, "xlsx")
    assert pieces({"sources": sources}) == {' فروش:row:1': '["۱۲۳"]', 'فروش:row:1': '["۴۵۶"]'}


def test_expanded_archive_limit_before_workbook_parser(tmp_path, monkeypatch):
    path = tmp_path / "source.xlsx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("oversize", b"x" * (25 * 1024 * 1024 + 1))
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *_args, **_kw: pytest.fail("must preflight"))
    with pytest.raises(DocumentError, match="resource_limit"):
        quick_extract.extract_units(path, "tesseract")


def test_total_projection_budget_rejects_instead_of_truncating():
    units = [ExtractedUnit(locator=f"CSV:row:{i}", text="[1]", method="csv", confidence=1) for i in range(800)]
    with pytest.raises(DocumentError, match="resource_limit"):
        pack_sources(units, "b" * 64, "csv")


@pytest.mark.parametrize("boundary", ["cells", "text", "sheets", "image-area", "image-type"])
def test_remaining_parser_limits_and_type_mismatch(tmp_path, monkeypatch, boundary):
    path = tmp_path / "source.csv"
    if boundary == "cells":
        # Keep text under a relaxed ceiling to isolate the independent cell budget.
        monkeypatch.setattr(quick_extract, "MAX_TEXT", 1024 * 1024)
        path.write_text((",".join(["1"] * 50) + "\n") * 401)
    elif boundary == "text":
        path.write_text(("1" * 5000 + "\n") * 9)
    elif boundary == "sheets":
        path = tmp_path / "source.xlsx"
        book = openpyxl.Workbook()
        for i in range(16):
            book.create_sheet(str(i))
        book.save(path)
        book.close()
    else:
        path = tmp_path / "source.png"
        image = Image.new("1", (5000, 5000)) if boundary == "image-area" else Image.new("RGB", (1, 1))
        image.save(path, format="BMP" if boundary == "image-type" else "PNG")
        if boundary == "image-area":
            monkeypatch.setattr(quick_extract.shutil, "which", lambda _: pytest.fail("no OCR before limit"))
    code = "invalid_document" if boundary == "image-type" else "resource_limit"
    with pytest.raises(DocumentError, match=code):
        quick_extract.extract_units(path, "tesseract")


def test_array_formula_is_rejected_even_without_string_value(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula
    book = openpyxl.Workbook()
    book.active["A1"] = ArrayFormula(ref="A1:A2", text="=ROW(A1:A2)")
    path = tmp_path / "source.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(DocumentError, match="unsupported_document"):
        quick_extract.extract_units(path, "tesseract")


def test_temporary_timeout_and_broken_result_cleanup(tmp_path, monkeypatch):
    def timeout(command, folder, seconds):
        (folder / "scratch").write_bytes(b"private")
        raise subprocess.TimeoutExpired(command, seconds)

    monkeypatch.setattr(quick_documents, "run_contained", timeout)
    with pytest.raises(DocumentError, match="timeout"):
        extract_temporary(CSV, "csv", temporary_parent=tmp_path)
    assert list(tmp_path.iterdir()) == []

    def broken(command, folder, seconds):
        (folder / "result.json").write_text("private invalid result")

    monkeypatch.setattr(quick_documents, "run_contained", broken)
    with pytest.raises(DocumentError, match="parser_failed"):
        extract_temporary(CSV, "csv", temporary_parent=tmp_path)
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("timeout", [False, True])
def test_process_tree_terminated_on_success_and_timeout(tmp_path, timeout):
    child = tmp_path / "child.py"
    child.write_text("import time\nfrom pathlib import Path\nwhile True:\n Path('beat').write_text('x')\n time.sleep(.03)\n")
    parent = tmp_path / "parent.py"
    parent.write_text(
        "import sys\nassert sys.stdin.buffer.read(1)==b'G'\n"
        "import subprocess,time,os\nfrom pathlib import Path\n"
        "assert 'HERMAN_PRIVATE_TEST' not in os.environ\n"
        "p=subprocess.Popen([sys.executable,'child.py'])\nPath('pid').write_text(str(p.pid))\n"
        "while not Path('beat').exists(): time.sleep(.01)\n"
        + ("time.sleep(60)\n" if timeout else "")
    )
    command = [sys.executable, str(parent)]
    if timeout:
        with pytest.raises(subprocess.TimeoutExpired):
            run_contained(command, tmp_path, 3)
    else:
        run_contained(command, tmp_path, 10)
    assert (tmp_path / "pid").exists()
    first = (tmp_path / "beat").stat().st_mtime_ns
    time.sleep(.15)
    assert (tmp_path / "beat").stat().st_mtime_ns == first
    # Removing the directory contents also checks for lingering Windows file handles.
    for path in tmp_path.iterdir():
        path.unlink()


def test_environment_is_allowlisted_and_scratch_is_local(tmp_path, monkeypatch):
    monkeypatch.setenv("HERMAN_PRIVATE_TEST", "not-forwarded")
    script = tmp_path / "probe.py"
    script.write_text("import sys\nassert sys.stdin.buffer.read(1)==b'G'\nimport os,json\n"
                      "from pathlib import Path\nPath('env.json').write_text(json.dumps(dict(os.environ)))\n")
    run_contained([sys.executable, str(script)], tmp_path, 10)
    environment = json.loads((tmp_path / "env.json").read_text())
    assert "HERMAN_PRIVATE_TEST" not in environment
    assert environment["TMP"] == environment["TEMP"] == environment["TMPDIR"] == str(tmp_path)
    assert "PYTHONPATH" not in environment


@pytest.mark.skipif(os.name != "nt", reason="Windows-specific Job Object assignment")
def test_failed_containment_never_releases_parser_handshake(tmp_path, monkeypatch):
    script = tmp_path / "probe.py"
    script.write_text("import sys\nassert sys.stdin.buffer.read(1)==b'G'\n"
                      "from pathlib import Path\nPath('should-not-exist').touch()\n")
    def fail(*_args):
        raise OSError("assignment refused")
    monkeypatch.setattr(process_guard.WindowsJob, "assign", fail)
    with pytest.raises(OSError, match="assignment refused"):
        run_contained([sys.executable, str(script)], tmp_path, 10)
    assert not (tmp_path / "should-not-exist").exists()


def request_stream(receive, headers=()):
    return Request({"type": "http", "method": "POST", "path": "/quick/documents/extract",
                    "headers": list(headers)}, receive)


def test_upload_stream_bound_and_slot_recovery(monkeypatch):
    async def scenario():
        slot = threading.BoundedSemaphore(1)
        async def receive():
            return {"type": "http.request", "body": b"x" * (MAX_UPLOAD + 1), "more_body": False}
        with pytest.raises(DocumentError, match="resource_limit"):
            await receive_and_extract(request_stream(receive), "csv", "tesseract", slot)
        assert slot.acquire(blocking=False)
        slot.release()
    asyncio.run(scenario())


def test_receive_timeout_and_disconnect_release_slot(monkeypatch):
    monkeypatch.setattr(quick_document_api, "RECEIVE_SECONDS", .02)
    async def scenario():
        slot = threading.BoundedSemaphore(1)
        async def slow():
            await asyncio.sleep(1)
        with pytest.raises(DocumentError, match="receive_timeout"):
            await receive_and_extract(request_stream(slow), "csv", "tesseract", slot)
        async def disconnected():
            return {"type": "http.disconnect"}
        with pytest.raises(DocumentError, match="invalid_document"):
            await receive_and_extract(request_stream(disconnected), "csv", "tesseract", slot)
        assert slot.acquire(blocking=False)
    asyncio.run(scenario())


def test_cancelled_request_keeps_slot_until_worker_finishes(monkeypatch):
    started, finish = threading.Event(), threading.Event()
    def work(*_args):
        started.set()
        assert finish.wait(5)
        return {}
    monkeypatch.setattr(quick_document_api, "extract_temporary", work)
    async def scenario():
        slot = threading.BoundedSemaphore(1)
        async def receive():
            return {"type": "http.request", "body": CSV, "more_body": False}
        task = asyncio.create_task(receive_and_extract(request_stream(receive), "csv", "tesseract", slot))
        while not started.is_set():
            await asyncio.sleep(.01)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        with pytest.raises(DocumentError, match="busy"):
            await receive_and_extract(request_stream(receive), "csv", "tesseract", slot)
        finish.set()
        for _ in range(200):
            if slot.acquire(blocking=False):
                slot.release()
                break
            await asyncio.sleep(.01)
        else:
            pytest.fail("cleanup did not release slot")
    try:
        asyncio.run(scenario())
    finally:
        finish.set()


def test_http_intake_context_and_draft_are_stateless(backend, monkeypatch):
    b = backend
    monkeypatch.setattr(b["db"], "transaction", lambda *_a, **_kw: pytest.fail("no business access"))
    captured = []
    class Draft:
        def advise(self, context):
            captured.append(context)
            return {"text_fa": "مقدار و واحد را بازبینی کنید", "model": "synthetic"}
    headers = {"Authorization": "Bearer " + b["tokens"][0]}
    before = set(b["blob_root"].rglob("*"))
    with TestClient(create_app(b["db"], b["blob_root"], Draft())) as client:
        denied = client.post("/quick/documents/extract?format=csv", content=CSV)
        assert denied.status_code == 401 and denied.headers["cache-control"] == "no-store"
        bad = client.post("/quick/documents/extract?format=private-secret", headers=headers, content=CSV)
        assert bad.status_code == 422 and "private-secret" not in bad.text
        assert bad.headers["cache-control"] == "no-store"
        result = client.post("/quick/documents/extract?format=csv", headers=headers, content=CSV)
        assert result.status_code == 200, result.text
        assert result.headers["cache-control"] == "no-store"
        body = {"question_fa": "این سند را بررسی کن", "sources": result.json()["sources"]}
        draft = client.post("/quick/advisor/draft", headers=headers, json=body)
        assert draft.status_code == 200, draft.text
        assert captured[0]["financial"] is None and captured[0]["history_used"] is False
        assert captured[0]["supplied_sources"][0]["origin"]["file_sha256"] == hashlib.sha256(CSV).hexdigest()
        body["sources"][0]["text_fa"] = "private-document-body"
        invalid = client.post("/quick/advisor/context", headers=headers, json=body)
        assert invalid.status_code == 422 and invalid.headers["cache-control"] == "no-store"
        assert "private-document-body" not in invalid.text and "origin" not in invalid.text
        second = client.post("/quick/advisor/context", headers=headers, json={"question_fa": "سؤال بعدی"})
        assert second.status_code == 200 and second.json()["supplied_sources"] == []
    assert set(b["blob_root"].rglob("*")) == before


def test_revoked_credential_cannot_receive_extraction_result(backend, monkeypatch):
    b = backend
    original, calls = b["db"].authenticate, []
    def authenticate(token):
        calls.append(1)
        if len(calls) > 1:
            raise AccessDenied("revoked")
        return original(token)
    monkeypatch.setattr(b["db"], "authenticate", authenticate)
    monkeypatch.setattr(quick_document_api, "extract_temporary", lambda *_: {"sources": ["private"]})
    with TestClient(create_app(b["db"], b["blob_root"])) as client:
        result = client.post("/quick/documents/extract?format=csv", content=CSV,
                             headers={"Authorization": "Bearer " + b["tokens"][0]})
        assert result.status_code == 401 and "private" not in result.text
        assert result.headers["cache-control"] == "no-store"


def test_intake_http_failures_keep_no_store_and_generic_body(backend, monkeypatch):
    b = backend
    headers = {"Authorization": "Bearer " + b["tokens"][0]}
    with TestClient(create_app(b["db"], b["blob_root"])) as client:
        for code, status in [("parser_failed", 503), ("timeout", 504), ("busy", 503)]:
            def failed(*_args):
                raise DocumentError(code)
            monkeypatch.setattr(quick_document_api, "extract_temporary", failed)
            response = client.post("/quick/documents/extract?format=csv", headers=headers, content=CSV)
            assert response.status_code == status and response.headers["cache-control"] == "no-store"
            assert response.json()["code"] == code
            if code == "busy":
                assert response.headers["retry-after"] == "5"
        for length, status in [(str(MAX_UPLOAD + 1), 413), ("9" * 5000, 422), ("-1", 422)]:
            response = client.post("/quick/documents/extract?format=csv", content=CSV,
                                   headers=headers | {"Content-Length": length})
            assert response.status_code == status and response.headers["cache-control"] == "no-store"
